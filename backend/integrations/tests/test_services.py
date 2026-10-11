import json
import tempfile
from datetime import date, datetime, timedelta, timezone as datetime_timezone
from io import BytesIO
from unittest.mock import patch

import httpx
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document as WordDocument
from openpyxl import load_workbook
from rest_framework.test import APIClient

from planner.models import Attendee, AuditEvent, Document, FollowUp, Issue, Meeting, PBCRequest, Project, TeamMember, TeamUpdate
from integrations.ai import capabilities, generate_draft
from integrations.connectors import ConnectorError, GitLab, Jira, fetch_issues, validate_base_url
from integrations.exchange import apply_import, parse_workbook, plan_import, safe_office_bytes, workbook_bytes
from integrations.models import AIDraft, Connection, IntegrationJob, RemoteIssue, RemoteSnapshot
from integrations.tasks import perform_draft, perform_sync
from integrations.word import extract_meeting, meeting_template


class IntegrationTestCase(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="integrations", password="very-strong-test-password")
        self.other = get_user_model().objects.create_user(username="another", password="very-strong-test-password")
        self.project = Project.objects.create(workspace=self.user, name="Portal", client_key="legacy-project")
        self.other_project = Project.objects.create(workspace=self.other, name="Other portal")
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def connection(self, provider="gitlab", resource="org/repo"):
        item = Connection.objects.create(workspace=self.user, project=self.project, name="Board", provider=provider, base_url="https://boards.example", resource=resource)
        item.set_token("server-secret-token")
        item.save()
        return item


class ExchangeTests(IntegrationTestCase):
    def workbook(self):
        self.pbc = PBCRequest.objects.create(workspace=self.user, project=self.project, request="Sample data", owner_name="Alex", request_date=date(2026, 10, 1), due_date=date(2026, 10, 15), comment="=literal formula-like text")
        return load_workbook(BytesIO(workbook_bytes("pbc", self.project)))

    def bytes(self, workbook):
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def test_populated_export_round_trip_and_formula_protection(self):
        workbook = self.workbook()
        self.assertEqual(workbook["PBC"]["A2"].value, self.pbc.reference)
        self.assertEqual(workbook["PBC"]["G2"].data_type, "s")
        rows = parse_workbook("pbc", self.project, self.bytes(workbook))
        preview = plan_import("pbc", self.project, rows)
        self.assertEqual(preview[0]["action"], "unchanged")
        self.assertFalse(preview[0]["errors"])

    def test_dates_require_reason_and_upsert_is_atomic(self):
        workbook = self.workbook()
        workbook["PBC"]["E2"] = date(2026, 10, 20)
        workbook["PBC"].append([None, "", "Alex", date(2026, 10, 1), None, "Open", ""])
        rows = parse_workbook("pbc", self.project, self.bytes(workbook))
        with self.assertRaisesMessage(ValueError, "Correct all"):
            apply_import("pbc", self.project, rows, "Client requested a later date")
        self.pbc.refresh_from_db()
        self.assertEqual(self.pbc.due_date, date(2026, 10, 15))
        workbook["PBC"].delete_rows(3)
        rows = parse_workbook("pbc", self.project, self.bytes(workbook))
        with self.assertRaisesMessage(ValueError, "Enter a reason"):
            apply_import("pbc", self.project, rows)
        changes = apply_import("pbc", self.project, rows, "Client requested a later date")
        self.assertEqual(len(changes), 1)
        self.pbc.refresh_from_db()
        self.assertEqual(self.pbc.due_date, date(2026, 10, 20))
        event = AuditEvent.objects.filter(action="excel_import").get()
        self.assertEqual(event.before["due_date"], "2026-10-15")
        self.assertEqual(event.after["due_date"], "2026-10-20")
        self.assertEqual(event.reason, "Client requested a later date")

    def test_blank_id_adds_unknown_id_or_cross_project_fails(self):
        workbook = self.workbook()
        workbook["PBC"].append([None, "Clarify rules", "Jamie", date(2026, 10, 3), None, "Open", ""])
        rows = parse_workbook("pbc", self.project, self.bytes(workbook))
        apply_import("pbc", self.project, rows)
        self.assertEqual(PBCRequest.objects.filter(project=self.project).count(), 2)
        self.assertEqual(PBCRequest.objects.filter(project=self.project).values("reference").distinct().count(), 2)
        with self.assertRaisesMessage(ValueError, "another project"):
            parse_workbook("pbc", self.other_project, self.bytes(workbook))
        workbook["PBC"]["A2"] = "PBC-FOREIGN"
        preview = plan_import("pbc", self.project, parse_workbook("pbc", self.project, self.bytes(workbook)))
        self.assertTrue(preview[0]["errors"])

    def test_formula_duplicate_and_wrong_date_rejected(self):
        workbook = self.workbook()
        workbook["PBC"]["B2"] = '=HYPERLINK("https://example.com")'
        workbook["PBC"]["E2"] = "not-a-date"
        workbook["PBC"].append([self.pbc.reference, "Duplicate", "Alex", None, None, "Open", ""])
        preview = plan_import("pbc", self.project, parse_workbook("pbc", self.project, self.bytes(workbook)))
        self.assertTrue(preview[0]["errors"])
        self.assertIn("Duplicate ID in this workbook.", preview[1]["errors"])

    def test_unassigned_owner_and_free_text_whitespace_roundtrip(self):
        PBCRequest.objects.create(workspace=self.user, project=self.project, request="Await owner", owner_name="", comment="  First paragraph\n\nSecond paragraph  \n")
        rows = parse_workbook("pbc", self.project, workbook_bytes("pbc", self.project))
        self.assertEqual(rows[0]["data"]["comment"], "  First paragraph\n\nSecond paragraph  \n")
        preview = plan_import("pbc", self.project, rows)
        self.assertFalse(preview[0]["errors"])
        self.assertEqual(preview[0]["action"], "unchanged")

    def test_api_preview_does_not_mutate_and_enforces_workspace(self):
        workbook = self.workbook()
        workbook["PBC"]["B2"] = "Updated request"
        upload = SimpleUploadedFile("pbc.xlsx", self.bytes(workbook))
        response = self.client.post(f"/api/integrations/projects/{self.project.client_key}/pbc/import/", {"file": upload}, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["can_import"])
        self.pbc.refresh_from_db()
        self.assertEqual(self.pbc.request, "Sample data")
        response = self.client.get(f"/api/integrations/projects/{self.other_project.pk}/pbc/template.xlsx")
        self.assertEqual(response.status_code, 404)

    def test_issue_template_updates_and_keeps_missing_rows(self):
        issue = Issue.objects.create(workspace=self.user, project=self.project, title="Missing report", raised_date=date(2026, 10, 1), owner_text="Minh", notes="Initial investigation")
        untouched = Issue.objects.create(workspace=self.user, project=self.project, title="Untouched", raised_date=date(2026, 10, 1))
        workbook = load_workbook(BytesIO(workbook_bytes("issues", self.project)))
        sheet = workbook["Issues"]
        sheet["D2"] = "Resolved"
        sheet["J2"] = "Report restored"
        sheet.delete_rows(3)
        rows = parse_workbook("issues", self.project, self.bytes(workbook))
        result = apply_import("issues", self.project, rows)
        self.assertEqual(len(result), 1)
        issue.refresh_from_db()
        self.assertEqual(issue.status, "Resolved")
        self.assertEqual(issue.resolution, "Report restored")
        self.assertTrue(Issue.objects.filter(pk=untouched.pk).exists())

    def test_issue_owners_export_and_upsert_keeps_assignment_consistent(self):
        old_owner = TeamMember.objects.create(workspace=self.user, name="Minh")
        new_owner = TeamMember.objects.create(workspace=self.user, name="Sarah")
        issue = Issue.objects.create(workspace=self.user, project=self.project, title="Assigned issue", raised_date=date(2026, 10, 1))
        issue.owners.add(old_owner)
        workbook = load_workbook(BytesIO(workbook_bytes("issues", self.project)))
        self.assertEqual(workbook["Issues"]["E2"].value, "Minh")
        rows = parse_workbook("issues", self.project, self.bytes(workbook))
        self.assertEqual(plan_import("issues", self.project, rows)[0]["action"], "unchanged")
        workbook["Issues"]["E2"] = "Sarah"
        apply_import("issues", self.project, parse_workbook("issues", self.project, self.bytes(workbook)))
        issue.refresh_from_db()
        self.assertEqual(list(issue.owners.all()), [new_owner])
        self.assertEqual(issue.owner_text, "Sarah")


class ConnectorTests(IntegrationTestCase):
    def test_credentials_encrypted_write_only_and_clear_on_server_change(self):
        connection = self.connection()
        self.assertNotIn("server-secret-token", connection.encrypted_token)
        self.assertEqual(connection.get_token(), "server-secret-token")
        response = self.client.get(f"/api/integrations/connections/{connection.pk}/")
        self.assertNotIn("server-secret-token", response.content.decode())
        self.assertTrue(response.data["has_credential"])
        response = self.client.patch(f"/api/integrations/connections/{connection.pk}/", {"base_url": "https://different.example"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        connection.refresh_from_db()
        self.assertEqual(connection.get_token(), "")
        response = self.client.patch(f"/api/integrations/connections/{connection.pk}/", {"project": str(self.other_project.pk)}, format="json")
        self.assertEqual(response.status_code, 404)

    def test_gitlab_pagination_get_only_and_encoded_repo_path(self):
        requests = []
        def respond(request):
            requests.append(request)
            page = request.url.params.get("page")
            body = [{"id": int(page), "iid": int(page), "title": "Review API", "state": "opened", "assignees": [], "updated_at": "2026-10-01"}]
            return httpx.Response(200, json=body, headers={"X-Next-Page": "2" if page == "1" else ""})
        records = fetch_issues(self.connection(), httpx.MockTransport(respond))
        self.assertEqual(len(records), 2)
        self.assertTrue(all(r.method == "GET" for r in requests))
        self.assertIn("org%2Frepo", str(requests[0].url))

    def test_redirect_not_followed_or_token_leaked(self):
        requests = []
        def respond(request):
            requests.append(request)
            return httpx.Response(302, headers={"Location": "https://hostile.example"})
        with self.assertRaises(ConnectorError):
            fetch_issues(self.connection(), httpx.MockTransport(respond))
        self.assertEqual(len(requests), 1)

    def test_jira_server_page_cap_honoured(self):
        offsets = []
        def respond(request):
            start = int(request.url.params["startAt"])
            offsets.append(start)
            rows = [{"id": str(i), "key": f"PROJ-{i}", "fields": {"summary": "Example", "status": {"name": "Done"}}} for i in range(start, min(start + 2, 5))]
            return httpx.Response(200, json={"issues": rows, "total": 5})
        records = Jira("https://jira.example", "12", "token", transport=httpx.MockTransport(respond)).issues()
        self.assertEqual(offsets, [0, 2, 4])
        self.assertEqual(len(records), 5)

    def test_sync_preserves_history_and_does_not_create_personal_tasks(self):
        connection = self.connection()
        row = {"id": 9, "iid": 2, "title": "Review API", "state": "opened", "assignees": []}
        transport = httpx.MockTransport(lambda r: httpx.Response(200, json=[row]))
        job = IntegrationJob.objects.create(workspace=self.user, project=self.project, connection=connection, kind="board_sync")
        perform_sync(job.pk, transport)
        job.refresh_from_db()
        self.assertEqual(job.status, "completed", job.error)
        self.assertEqual(RemoteIssue.objects.count(), 1)
        self.assertEqual(RemoteSnapshot.objects.count(), 1)
        perform_sync(job.pk, transport)
        self.assertEqual(RemoteSnapshot.objects.count(), 1)
        row["state"] = "closed"
        perform_sync(job.pk, transport)
        self.assertEqual(RemoteSnapshot.objects.count(), 2)
        self.assertEqual(self.project.activities.count(), 0)

    def test_discrepancies_use_today_accepted_updates_and_status_meaning(self):
        connection = self.connection()
        issue = RemoteIssue.objects.create(connection=connection, external_id="1", key="#12", title="API", status="opened")
        update = TeamUpdate.objects.create(workspace=self.user, project=self.project, update_date=timezone.localdate(), source_text="API done", team="Hanoi", accepted=True, proposed_changes=[{"key": "#12", "status": "Done"}])
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(len(response.data["discrepancies_today"]), 1)
        issue.status = "closed"
        issue.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(response.data["discrepancies_today"], [])
        issue.status = "opened"
        issue.save()
        update.update_date = timezone.localdate() - timedelta(days=1)
        update.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(response.data["discrepancies_today"], [])
        self.assertEqual(TeamUpdate.objects.get(pk=update.pk).source_text, "API done")

    def test_duplicate_gitlab_numbers_do_not_create_false_cross_repo_discrepancy(self):
        connection = self.connection()
        issue = RemoteIssue.objects.create(connection=connection, external_id="1", key="#12", title="API", status="opened")
        other = self.connection(resource="org/another")
        RemoteIssue.objects.create(connection=other, external_id="99", key="#12", title="Other task", status="opened")
        update = TeamUpdate.objects.create(workspace=self.user, project=self.project, update_date=timezone.localdate(), source_text="Done", accepted=True, proposed_changes=[{"key": "#12", "status": "Done"}])
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(response.data["discrepancies_today"], [])
        update.proposed_changes = [{"key": "org/repo#12", "status": "Done"}]
        update.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(len(response.data["discrepancies_today"]), 1)

    def test_gitlab_open_state_accepts_all_nonterminal_workflow_statuses(self):
        connection = self.connection()
        issue = RemoteIssue.objects.create(connection=connection, external_id="1", key="#12", title="API", status="opened")
        update = TeamUpdate.objects.create(workspace=self.user, project=self.project, update_date=timezone.localdate(), source_text="API is underway", team="Hanoi", accepted=True)
        for reported in ["In progress", "Blocked", "Awaiting my review", "Not started"]:
            with self.subTest(reported=reported):
                update.proposed_changes = [{"key": "org/repo#12", "status": reported}]
                update.save()
                response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
                self.assertEqual(response.data["discrepancies_today"], [])
        update.proposed_changes = [{"key": "org/repo#12", "status": "Done"}]
        update.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(len(response.data["discrepancies_today"]), 1)
        issue.status = "closed"
        issue.save()
        update.proposed_changes = [{"key": "org/repo#12", "status": "In progress"}]
        update.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(len(response.data["discrepancies_today"]), 1)

    def test_jira_rich_workflow_statuses_still_compare_directly(self):
        connection = self.connection(provider="jira", resource="12")
        issue = RemoteIssue.objects.create(connection=connection, external_id="1", key="APP-12", title="API", status="To do")
        TeamUpdate.objects.create(workspace=self.user, project=self.project, update_date=timezone.localdate(), source_text="API in progress", team="Hanoi", accepted=True, proposed_changes=[{"key": "APP-12", "status": "In progress"}])
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(len(response.data["discrepancies_today"]), 1)
        issue.status = "In progress"
        issue.save()
        response = self.client.get(f"/api/integrations/remote-issues/{issue.pk}/")
        self.assertEqual(response.data["discrepancies_today"], [])

    def test_connection_edit_during_pull_rejects_old_source_response(self):
        connection = self.connection()
        job = IntegrationJob.objects.create(workspace=self.user, project=self.project, connection=connection, kind="board_sync")
        def respond(request):
            Connection.objects.filter(pk=connection.pk).update(resource="org/different")
            return httpx.Response(200, json=[{"id": 1, "iid": 1, "title": "Old repository item", "state": "opened"}])
        perform_sync(job.pk, httpx.MockTransport(respond))
        job.refresh_from_db()
        self.assertEqual(job.status, "failed")
        self.assertIn("changed while", job.error)
        self.assertEqual(RemoteIssue.objects.count(), 0)


class WordAndAITests(IntegrationTestCase):
    def test_template_prefills_attendees_pending_actions_and_singapore_time(self):
        meeting = Meeting.objects.create(workspace=self.user, project=self.project, title="Walkthrough", starts_at=datetime(2026, 10, 12, 6, tzinfo=datetime_timezone.utc))
        Attendee.objects.create(workspace=self.user, meeting=meeting, name="Alex", email="alex@example.com")
        FollowUp.objects.create(workspace=self.user, project=self.project, title="Confirm revised scope")
        FollowUp.objects.create(workspace=self.user, project=self.project, title="Already complete", status="Done")
        PBCRequest.objects.create(workspace=self.user, project=self.project, request="Supply sample data")
        document = WordDocument(BytesIO(meeting_template(meeting, self.project)))
        text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("12 Oct 2026, 14:00", text)
        self.assertIn("Alex", text)
        self.assertIn("Confirm revised scope", text)
        self.assertIn("Supply sample data", text)
        self.assertNotIn("Already complete", text)

    def test_docx_template_and_extraction_retain_structured_attendees(self):
        content = meeting_template(project=self.project)
        word = WordDocument(BytesIO(content))
        attendees = word.tables[1]
        for cell, value in zip(attendees.rows[1].cells, ["Alex", "Client A", "Sponsor", "alex@example.com", "Attended"]):
            cell.text = value
        word.add_heading("Discussion", level=1)
        word.add_paragraph("Client asked to review the export format. No agreement was confirmed.")
        output = BytesIO()
        word.save(output)
        result = extract_meeting(output.getvalue())
        self.assertEqual(result["attendees"][0]["name"], "Alex")
        self.assertIn("No agreement", result["text"])
        self.assertNotIn("List the topics to cover.", result["text"])

    @override_settings(OLLAMA_BASE_URL="http://ollama:11434", OLLAMA_MODEL="test-model")
    def test_ollama_drafts_have_no_side_effects(self):
        sent = []
        output = {"summary": "Pending client confirmation", "decisions": [], "actions": [], "pbc": [], "scope_changes": []}
        def respond(request):
            sent.append(json.loads(request.content))
            return httpx.Response(200, json={"message": {"content": json.dumps(output)}})
        draft = AIDraft.objects.create(workspace=self.user, project=self.project, kind="meeting", source_text="Discussed export. Client has not confirmed.")
        job = IntegrationJob.objects.create(workspace=self.user, project=self.project, kind="ai_draft")
        perform_draft(job.pk, draft.pk, httpx.MockTransport(respond))
        job.refresh_from_db()
        draft.refresh_from_db()
        self.assertEqual(job.status, "completed", job.error)
        self.assertEqual(draft.status, "draft")
        self.assertEqual(draft.structured, output)
        self.assertEqual(self.project.pbc_requests.count(), 0)
        self.assertIn("untrusted evidence", sent[0]["messages"][0]["content"])

    def test_file_upload_download_and_extraction_are_private(self):
        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder):
            content = meeting_template(project=self.project)
            upload = SimpleUploadedFile("notes.docx", content, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            response = self.client.post("/api/integrations/files/", {"file": upload, "context": f"project:{self.project.client_key}"}, format="multipart")
            self.assertEqual(response.status_code, 201, response.data)
            file_id, url = response.data["id"], response.data["url"]
            response = self.client.post("/api/integrations/meetings/import/", {"file_id": file_id}, format="json")
            self.assertEqual(response.status_code, 200, response.data)
            self.assertIn("text", response.data)
            self.assertEqual(Meeting.objects.count(), 0)
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(b"".join(response.streaming_content), content)
            self.assertIn("no-store", response["Cache-Control"])
            self.client.force_authenticate(self.other)
            self.assertEqual(self.client.get(url).status_code, 404)
            self.client.force_authenticate(None)
            self.assertIn(self.client.get(url).status_code, [401, 403])

    def test_stale_upload_rejected_before_storage_write(self):
        from workspace.models import WorkspaceState
        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder):
            state = WorkspaceState.objects.get(user=self.user)
            revision = state.revision
            upload = SimpleUploadedFile("notes.txt", b"Notes")
            response = self.client.post("/api/integrations/files/", {"file": upload, "project": str(self.project.pk), "expected_revision": revision - 1}, format="multipart")
            self.assertEqual(response.status_code, 409)
            self.assertEqual(Document.objects.count(), 0)
            upload = SimpleUploadedFile("notes.txt", b"Notes")
            response = self.client.post("/api/integrations/files/", {"file": upload, "project": str(self.project.pk), "expected_revision": revision}, format="multipart")
            self.assertEqual(response.status_code, 201, response.data)
            self.assertEqual(response.data["previous_revision"], revision)
            self.assertEqual(response.data["workspace_revision"], revision + 1)

    def test_real_session_upload_handles_lazy_authenticated_user(self):
        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder):
            self.client.force_authenticate(None)
            self.assertTrue(self.client.login(username="integrations", password="very-strong-test-password"))
            response = self.client.post("/api/integrations/files/", {"file": SimpleUploadedFile("notes.txt", b"Notes"), "project": str(self.project.pk)}, format="multipart")
            self.assertEqual(response.status_code, 201, response.data)

    def test_word_commit_keeps_original_and_source_without_silent_actions(self):
        with tempfile.TemporaryDirectory() as folder, override_settings(MEDIA_ROOT=folder):
            content = meeting_template(project=self.project)
            upload = SimpleUploadedFile("minutes.docx", content)
            response = self.client.post("/api/integrations/meetings/import/", {"file": upload, "project": str(self.project.pk), "commit": "true", "title": "Client walkthrough", "starts_at": "2026-10-12T14:00:00+08:00"}, format="multipart")
            self.assertEqual(response.status_code, 201, response.data)
            meeting = Meeting.objects.get(pk=response.data["meeting"])
            document = Document.objects.get(pk=response.data["document"])
            self.assertEqual(document.file.read(), content)
            self.assertTrue(meeting.transcript)
            self.assertEqual(meeting.decisions.count(), 0)
            self.assertEqual(meeting.status, "draft")

    @patch("integrations.views.create_draft")
    def test_ai_endpoint_queues_and_returns_pollable_ids(self, enqueue):
        response = self.client.post("/api/integrations/ai/drafts/", {"kind": "email", "source_text": "Ask Alex for the requested data.", "project": self.project.client_key}, format="json")
        self.assertEqual(response.status_code, 202, response.data)
        enqueue.assert_called_once()
        job = self.client.get(f"/api/integrations/jobs/{response.data['job']}/")
        self.assertEqual(job.data["status"], "queued")
        draft = self.client.get(f"/api/integrations/ai/drafts/{response.data['draft']}/")
        self.assertEqual(draft.status_code, 200)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f"/api/integrations/jobs/{response.data['job']}/").status_code, 404)

    @override_settings(OLLAMA_BASE_URL="http://ollama:11434", OLLAMA_MODEL="local-model", CLOUD_AI_BASE_URL="https://cloud.example/v1", CLOUD_AI_MODEL="cloud-model", CLOUD_AI_API_KEY="private-cloud-key")
    def test_cloud_is_explicit_and_never_an_automatic_fallback(self):
        calls = []
        def respond(request):
            calls.append(request)
            if request.url.host == "ollama":
                return httpx.Response(503)
            return httpx.Response(200, json={"choices": [{"message": {"content": "Draft email"}}]})
        transport = httpx.MockTransport(respond)
        with self.assertRaises(ConnectorError):
            generate_draft("email", "Client has not replied.", transport)
        self.assertEqual([request.url.host for request in calls], ["ollama"])
        result = generate_draft("email", "Client has not replied.", transport, provider="cloud")
        self.assertEqual(result["content"], "Draft email")
        self.assertEqual(calls[-1].headers["Authorization"], "Bearer private-cloud-key")
        response = self.client.get("/api/integrations/ai/capabilities/")
        self.assertEqual(response.data["default_provider"], "ollama")
        self.assertTrue(response.data["providers"][1]["configured"])
        self.assertNotIn("private-cloud-key", response.content.decode())
