"""Populated Excel templates with validated, project-scoped, atomic upserts."""
from copy import copy
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from django.contrib.auth import get_user_model
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


MAX_UPLOAD = 20 * 1024 * 1024
MAX_EXPANDED = 80 * 1024 * 1024
MAX_ROWS = 1000
CONFIG = {
    "pbc": {
        "model": "PBCRequest", "sheet": "PBC",
        "headers": ["PBC ID", "Request", "Owner", "Request date", "Due date", "Status", "Comment"],
        "fields": ["reference", "request", "owner_name", "request_date", "due_date", "status", "comment"],
        "dates": ["request_date", "due_date"], "required": ["request"],
        "free_text": ["comment"],
        "choices": {"status": ["Open", "In progress", "Waiting", "Done"]},
        "defaults": {"status": "Open"}, "fingerprint": ["request", "owner_name", "request_date"],
    },
    "issues": {
        "model": "Issue", "sheet": "Issues",
        "headers": ["Issue ID", "Issue", "Severity", "Status", "Owners", "Raised date", "Due date", "Description", "Notes", "Resolution"],
        "fields": ["reference", "title", "severity", "status", "owner_text", "raised_date", "due_date", "description", "notes", "resolution"],
        "dates": ["raised_date", "due_date"], "required": ["title", "raised_date"],
        "free_text": ["description", "notes", "resolution"],
        "choices": {"status": ["Open", "In progress", "Blocked", "Resolved", "Closed"], "severity": ["Low", "Medium", "High", "Critical"]},
        "defaults": {"status": "Open", "severity": "Medium"}, "fingerprint": ["title", "raised_date", "owner_text"],
    },
}


def safe_office_bytes(upload, suffix):
    if not upload.name.lower().endswith(suffix):
        raise ValueError(f"Upload a {suffix} file.")
    if getattr(upload, "size", 0) > MAX_UPLOAD:
        raise ValueError("The upload exceeds 20 MB.")
    content = upload.read(MAX_UPLOAD + 1)
    if len(content) > MAX_UPLOAD:
        raise ValueError("The upload exceeds 20 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > 5000 or sum(item.file_size for item in files) > MAX_EXPANDED:
                raise ValueError("The document expands beyond the supported size. Split it into smaller files.")
            if any(item.flag_bits & 1 for item in files):
                raise ValueError("Password-protected documents are not supported.")
    except BadZipFile as exc:
        raise ValueError("This is not a valid Office document.") from exc
    return content


def config_and_model(kind):
    from planner import models
    if kind not in CONFIG:
        raise ValueError("Choose PBC or issues for Excel exchange.")
    config = CONFIG[kind]
    return config, getattr(models, config["model"])


def text_cell(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"  # Never turn a user's text beginning '=' into an Excel formula.
    return cell


def exchange_value(instance, field):
    value = getattr(instance, field)
    if field == "owner_text" and not value and hasattr(instance, "owners") and not instance._state.adding:
        return ", ".join(instance.owners.order_by("name").values_list("name", flat=True))
    return value


def workbook_bytes(kind, project):
    config, model = config_and_model(kind)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config["sheet"]
    sheet.append(config["headers"])
    sheet.freeze_panes = "B2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A6E")
    for row_index, item in enumerate(model.objects.filter(project=project, workspace=project.workspace).order_by("reference"), 2):
        for column, field in enumerate(config["fields"], 1):
            value = exchange_value(item, field)
            cell = text_cell(sheet, row_index, column, value)
            if field in config["dates"]:
                cell.number_format = "yyyy-mm-dd"
    for index, field in enumerate(config["fields"], 1):
        sheet.column_dimensions[get_column_letter(index)].width = 40 if field in ("request", "title", "comment", "description", "notes", "resolution") else 22
        if field in config["choices"]:
            validator = DataValidation(type="list", formula1='"' + ",".join(config["choices"][field]) + '"')
            validator.error = "Choose a value from the list."
            validator.errorTitle = "Invalid value"
            validator.showErrorMessage = True
            sheet.add_data_validation(validator)
            validator.add(f"{get_column_letter(index)}2:{get_column_letter(index)}1001")
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(config['fields']))}{max(sheet.max_row, 2)}"
    metadata = workbook.create_sheet("_Planner")
    metadata.append(["Project ID", str(project.pk)])
    metadata.append(["Project reference", project.reference])
    metadata.append(["Kind", kind])
    metadata.append(["Instructions", "Keep existing IDs to update. Leave ID blank for new rows. Missing rows are never deleted."])
    metadata.sheet_state = "hidden"
    instructions = workbook.create_sheet("Instructions")
    for row in [["How to use this populated template"], ["Keep IDs for updates. Leave the ID blank to add a new item."], ["Do not rename headers or remove the hidden project information."], ["Removing a row here does not delete an item from the app."], ["Dates: yyyy-mm-dd. Use plain values; formulas are not imported."], ["Upload for a preview, resolve all errors, then confirm the import."], ["When dates change, enter the reason in the app for accountability."]]:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 100
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def parse_date(value):
    if value in ("", None):
        return None
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            raise ValueError("Use a date without a time.")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise ValueError("Use yyyy-mm-dd or a formatted Excel date.")


def primitive(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def parse_workbook(kind, project, content):
    config, _ = config_and_model(kind)
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("This workbook cannot be read. Download a fresh template and try again.") from exc
    try:
        if config["sheet"] not in workbook:
            raise ValueError(f"Keep the worksheet named {config['sheet']}.")
        if "_Planner" not in workbook or str(workbook["_Planner"]["B1"].value) != str(project.pk) or workbook["_Planner"]["B3"].value != kind:
            raise ValueError("This template belongs to another project or is missing its project information. Download this project's template.")
        sheet = workbook[config["sheet"]]
        if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 100:
            raise ValueError("Upload at most 1,000 rows at a time.")
        rows = sheet.iter_rows(max_col=len(config["fields"]))
        headers = next(rows, ())
        if [str(cell.value or "").strip().lower() for cell in headers] != [h.lower() for h in config["headers"]]:
            raise ValueError("Keep all template column headers unchanged.")
        parsed = []
        for row_number, row in enumerate(rows, 2):
            if all(cell.value in (None, "") for cell in row):
                continue
            data, errors = {}, []
            for field, cell in zip(config["fields"], row):
                if cell.data_type in ("f", "e"):
                    errors.append(f"{field}: Use a value, not a formula or Excel error.")
                try:
                    value = str(cell.value if cell.value is not None else "")
                    data[field] = parse_date(cell.value) if field in config["dates"] else value if field in config["free_text"] else value.strip()
                except ValueError as exc:
                    data[field] = None
                    errors.append(f"{field}: {exc}")
            for key, default in config["defaults"].items():
                data[key] = data[key] or default
            parsed.append({"row": row_number, "data": data, "errors": errors})
        return parsed
    finally:
        workbook.close()


def plan_import(kind, project, rows, lock=False):
    config, model = config_and_model(kind)
    queryset = model.objects.filter(project=project, workspace=project.workspace)
    if lock:
        queryset = queryset.select_for_update()
    existing = {item.reference: item for item in queryset}
    seen, new_fingerprints = set(), set()
    fingerprints = {tuple(str(exchange_value(item, f) or "").casefold() for f in config["fingerprint"]) for item in existing.values()}
    result = []
    for entry in rows:
        data, errors = entry["data"], list(entry["errors"])
        reference = data["reference"]
        old = existing.get(reference) if reference else None
        if reference:
            if reference in seen:
                errors.append("Duplicate ID in this workbook.")
            if not old:
                errors.append("ID does not belong to this project. Leave it blank for a new item.")
            seen.add(reference)
        else:
            fingerprint = tuple(str(data[f] or "").casefold() for f in config["fingerprint"])
            if fingerprint in fingerprints or fingerprint in new_fingerprints:
                errors.append("This item already exists. Use its ID to update it.")
            new_fingerprints.add(fingerprint)
        for field in config["required"]:
            if not data[field]:
                errors.append(f"{field} is required.")
        for field, choices in config["choices"].items():
            if data[field] not in choices:
                errors.append(f"{field}: Choose {', '.join(choices)}.")
        start = data[config["dates"][0]]
        if start and data["due_date"] and start > data["due_date"]:
            errors.append("Due date cannot be before the request/raised date.")
        instance = copy(old) if old else model(project=project, workspace=project.workspace)
        for field in config["fields"][1:]:
            setattr(instance, field, data[field])
        try:
            instance.full_clean(exclude=["reference"], validate_unique=False, validate_constraints=False)
        except ValidationError as exc:
            errors.extend(exc.messages)
        changes = [f for f in config["fields"][1:] if old and exchange_value(old, f) != data[f]]
        result.append({"row": entry["row"], "data": {k: primitive(v) for k, v in data.items()}, "errors": errors, "changes": changes, "date_changed": bool(set(changes) & set(config["dates"])), "action": ("update" if changes else "unchanged") if old else "add", "instance": instance})
    return result


def public_plan(plan):
    return [{key: value for key, value in entry.items() if key != "instance"} for entry in plan]


@transaction.atomic
def apply_import(kind, project, rows, reason=""):
    from planner.audit import record_event
    # Match the lock order used by workspace autosave and canonical CRUD.
    get_user_model().objects.select_for_update().get(pk=project.workspace_id)
    project.__class__.objects.select_for_update().get(pk=project.pk)
    plan = plan_import(kind, project, rows, lock=True)
    if any(entry["errors"] for entry in plan):
        raise ValueError("Correct all workbook errors before importing.")
    if any(entry["date_changed"] for entry in plan) and not reason.strip():
        raise ValueError("Enter a reason for the changed dates.")
    config, _ = config_and_model(kind)
    changed = []
    for entry in plan:
        if entry["action"] == "unchanged":
            continue
        instance = entry["instance"]
        previous = instance.__class__.objects.filter(pk=instance.pk).first() if entry["action"] == "update" else None
        before = {f: primitive(exchange_value(previous, f)) for f in config["fields"]} if previous else {}
        instance.save()
        if kind == "issues" and (not previous or "owner_text" in entry["changes"]):
            from planner.models import TeamMember
            members = []
            for name in filter(None, (part.strip() for part in re.split(r"[,;]", instance.owner_text))):
                matches = list(TeamMember.objects.filter(workspace=project.workspace, name__iexact=name)[:2])
                if len(matches) == 1:
                    members.append(matches[0])
            instance.owners.set(members)
        after = {f: primitive(exchange_value(instance, f)) for f in config["fields"]}
        record_event(project.workspace, instance, "excel_import", before, after, reason)
        changed.append({"id": str(instance.pk), "reference": instance.reference, "action": entry["action"]})
    return changed
