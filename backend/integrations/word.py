"""Structured meeting templates; embedded screenshots remain in the original document."""
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document as WordDocument
from docx.shared import Inches, Pt
from django.utils import timezone


def meeting_template(meeting=None, project=None):
    document = WordDocument()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    document.add_heading("Meeting minutes", 0)
    document.add_paragraph("Daily Project Planner · Internal working record")
    meta = document.add_table(rows=0, cols=2)
    meta.style = "Light Shading Accent 1"
    for label, value in [
        ("Meeting title", meeting.title if meeting else ""),
        ("Project", project.name if project else ""),
        ("Date & time (Singapore)", timezone.localtime(meeting.starts_at, ZoneInfo("Asia/Singapore")).strftime("%d %b %Y, %H:%M") if meeting else ""),
        ("Location / meeting link", meeting.location if meeting else ""),
        ("Purpose", str(meeting.extra.get("purpose", "")) if meeting else ""),
    ]:
        cells = meta.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_heading("Attendees", level=1)
    attendees = document.add_table(rows=1, cols=5)
    attendees.style = "Light Shading Accent 1"
    for cell, label in zip(attendees.rows[0].cells, ["Name", "Organisation", "Role", "Email", "Attendance"]):
        cell.text = label
    existing_attendees = list(meeting.attendees.all()) if meeting else []
    for person in existing_attendees:
        for cell, value in zip(attendees.add_row().cells, [person.name, person.organisation, person.role, person.email, person.attendance]):
            cell.text = value
    for _ in range(max(1, 3 - len(existing_attendees))):
        attendees.add_row()
    context = project or (meeting.project or meeting.opportunity if meeting else None)
    if context:
        outstanding = [
            ("Follow-up", item.title, item.owner_name, item.due_date, item.status)
            for item in context.follow_ups.exclude(status="Done").order_by("due_date")
        ] + [
            (item.reference, item.request, item.owner_name, item.due_date, item.status)
            for item in context.pbc_requests.exclude(status="Done").order_by("due_date")
        ]
        document.add_heading("Previous outstanding actions", level=1)
        if outstanding:
            pending = document.add_table(rows=1, cols=5)
            pending.style = "Light Shading Accent 1"
            for cell, value in zip(pending.rows[0].cells, ["Type / ID", "Action / request", "Owner", "Due date", "Status"]):
                cell.text = value
            for row in outstanding:
                for cell, value in zip(pending.add_row().cells, row):
                    cell.text = str(value or "")
        else:
            document.add_paragraph("No outstanding actions recorded.")
    for heading, hint in [
        ("Agenda", "List the topics to cover."),
        ("Discussion", "Capture what was discussed. Paste screenshots directly below the relevant topic."),
        ("Decisions & agreed items", "Record the wording, who agreed, and whether confirmation is explicit or still pending."),
        ("Actions & follow-ups", "For each action, record the task, owner, due date and next follow-up date."),
        ("PBC requests", "Record each document or clarification requested from the client, owner and due date."),
        ("Scope & deadline changes", "Record the previous value, proposed change, reason, and confirmation status."),
        ("Next meeting", "Record the next meeting or check-in, if agreed."),
    ]:
        document.add_heading(heading, level=1)
        paragraph = document.add_paragraph(hint)
        paragraph.style = "Caption"
        document.add_paragraph(meeting.notes if meeting and heading == "Discussion" else "")
    footer = section.footer.paragraphs[0]
    footer.text = "Internal record · Review extracted content before accepting actions or decisions."
    footer.style = "Caption"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def extract_meeting(content):
    try:
        document = WordDocument(BytesIO(content))
    except Exception as exc:
        raise ValueError("The Word document cannot be read. Upload a .docx document.") from exc
    sections, heading = {}, "General"
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            heading = text
        elif paragraph.style.name == "Caption":
            continue  # Template prompts are guidance, not meeting evidence.
        else:
            sections.setdefault(heading, []).append(text)
        lines.append(text)
    attendees, metadata, tables = [], {}, []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append(rows)
        if rows and rows[0][:2] == ["Name", "Organisation"]:
            for row in rows[1:]:
                if row[0]:
                    attendees.append(dict(zip(["name", "organisation", "role", "email", "attendance"], row)))
        elif rows and rows[0][0] == "Meeting title":
            metadata.update({row[0]: row[1] for row in rows if len(row) >= 2})
        for row in rows:
            if any(row):
                lines.append(" | ".join(row))
    text = "\n".join(lines)
    if len(text) > 250000:
        raise ValueError("This meeting record is too long. Split it into smaller documents.")
    return {"text": text, "sections": {key: "\n".join(value) for key, value in sections.items()}, "attendees": attendees, "metadata": metadata, "tables": tables, "image_count": len(document.inline_shapes), "notice": "Screenshots remain supporting evidence in the original Word file. AI extraction creates drafts for your review."}
